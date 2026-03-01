import os

# ================= DESKTOP AUTO PATH =================
DESKTOP = os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop")
if not os.path.exists(DESKTOP):
    DESKTOP = os.path.join(os.path.expanduser("~"), "Desktop")


def _full_path(path):
    if not path:
        return None

    # if user gives full path
    if ":" in path or "\\" in path:
        return path

    # otherwise → desktop
    return os.path.join(DESKTOP, path)


# ================= CREATE FILE =================
def create_file(path="", content=""):
    path = _full_path(path)

    try:
        # ---------- WORD ----------
        if path.endswith(".docx"):
            from docx import Document
            doc = Document()
            if content:
                doc.add_paragraph(content)
            doc.save(path)
            return f"Word file created: {path}"

        # ---------- PDF ----------
        elif path.endswith(".pdf"):
            from reportlab.pdfgen import canvas
            c = canvas.Canvas(path)
            if content:
                c.drawString(100, 750, content)
            c.save()
            return f"PDF created: {path}"

        # ---------- EXCEL ----------
        elif path.endswith(".xlsx"):
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            if content:
                ws["A1"] = content
            wb.save(path)
            return f"Excel file created: {path}"

        # ---------- NORMAL FILE ----------
        else:
            with open(path, "w", encoding="utf-8") as f:
                f.write(content if content else "")
            return f"File created: {path}"

    except Exception as e:
        return f"Error creating file: {e}"


# ================= READ FILE =================
def read_file(path=""):
    path = _full_path(path)

    if not os.path.exists(path):
        return "File not found."

    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except:
        return "Cannot read this file type."


# ================= WRITE FILE =================
def write_file(path="", content=""):
    path = _full_path(path)

    try:
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        return f"File updated: {path}"
    except Exception as e:
        return f"Error writing file: {e}"


# ================= DELETE =================
def delete_file(path=""):
    path = _full_path(path)

    if not os.path.exists(path):
        return "File not found."

    try:
        os.remove(path)
        return f"Deleted: {path}"
    except Exception as e:
        return f"Error deleting: {e}"


# ================= CREATE FOLDER =================
def create_folder(path=""):
    path = _full_path(path)

    try:
        os.makedirs(path, exist_ok=True)
        return f"Folder created: {path}"
    except Exception as e:
        return f"Error creating folder: {e}"


# ================= SUMMARIZE FILE =================
def summarize_file(path="", llm=None):
    path = _full_path(path)

    if not os.path.exists(path):
        return "File not found."

    try:
        # try reading text
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read(8000)

        if not text.strip():
            return "File empty."

        if llm:
            return llm(f"Summarize this file:\n\n{text}")

        return "LLM not available."

    except Exception as e:
        return f"Error summarizing: {e}"